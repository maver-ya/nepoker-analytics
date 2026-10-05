# -*- coding: utf-8 -*-
"""Собирает data/data.json для сайта из Excel-файлов НеПокера.

Запуск (из папки site):   python build_data.py
Источники:
  2026/<турнир>/*.xlsx            - таблицы турниров (лист 'Главная')
  Данные/Рейтинг ЦИФЕРБЛАТ.xlsx   - очки рейтинга по событиям (из очков восстанавливаются места)
  Данные/Непокер в Циферблате.xlsx - лист 'Итоги' (призёры, чтобы привязать столбцы рейтинга)

Новый турнир: добавить запись в TOURNAMENTS и запустить скрипт ещё раз.
Места восстанавливаются по формуле рейтинга: очки = sqrt(N*K)/sqrt(место), т.е. место = (макс.очки / очки)^2.
"""
import openpyxl, glob, json, os, re, sys, statistics as st
from openpyxl.utils import column_index_from_string as ci

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, '..'))
RATING = os.path.join(ROOT, 'Данные', 'Рейтинг ЦИФЕРБЛАТ.xlsx')
HIST = os.path.join(ROOT, 'Данные', 'Непокер в Циферблате.xlsx')

# final_col  - столбец с итогом отборочных (без базового бонуса)
# base       - базовый бонус финалиста
# trophy     - множители 'охотничьих трофеев' по числу посещённых дней; price_col - столбец с ценой головы
TOURNAMENTS = [
    dict(id='Q2', folder='Январь 2026 Q2', title='Миссис Дейзи Дак Q-2', short='Q-2', final_col='O', base=3000, ref='Миссис Дейзи Дак Q2'),
    dict(id='K2', folder='Январь-февраль 2026 К2', title='Мистер Дональд Дак К-2', short='К-2', final_col='AF', base=3000, ref='Дональд Дак К2',
         flag='Нестандартные правила: кредит и «сделка с дьяволом» — стек сравним с другими только приблизительно', nonstandard=True),
    dict(id='A2', folder='Апрель А2', title='Сезон охоты А-2', short='А-2', final_col='W', base=3000, ref='Сезон Охоты А2',
         trophy={1: .5, 2: 1, 3: 1.5, 4: 2}, price_col='Y'),
    dict(id='34', folder='Апрель-Май 2026 34', title='Вэлью Бет 3/4', short='3-4', final_col='AE', base=3000, ref='Вэлью Бет 3/4',
         flag='Отдельная валюта: финальные фишки выдаются за место в каждом дне', nonstandard=True),
    dict(id='35', folder='Июнь-июль 2026 35', title='Билли Джонсон 3-5', short='3-5', final_col='I', base=3000, ref='Билли Джонсон 3-5',
         flag='Только 2 отборочных дня, в таблице нет стартового стека дня', nonstandard=True),
    dict(id='36', folder='Июль-август 2026 36', title='Джимми Саммерфилд 3-6', short='3-6', final_col=None, base=3000, ref='Джимми Саммерфилд 3-6',
         flag='В таблице нет данных (файл — копия 6-7): есть только места из рейтинга', use_table=False),
    dict(id='67', folder='Август 2026 67', title='Six Seven 6-7', short='6-7', final_col=None, base=3000, ref='Six Seven 6-7',
         flag='В таблице нет финального стека'),
    dict(id='37', folder='Август-Сентябрь 2026 3-7', title='Джо Хашем 3-7', short='3-7', final_col='AE', base=3800, ref='Джо Хашем 3-7'),
    dict(id='38', folder='Сентябрь-Октябрь 2026 38', title='Джонатан Томайо 3-8', short='3-8', final_col='U', base=3800, ref='Джонатан Томайо 3-8',
         trophy={1: 1, 2: 1.5, 3: 2}, price_col='W',
         note='Цены голов в Excel и на сайте турнира расходятся: Богдан А 6125 (Excel) / 6725 (сайт), Артём SUB 1765 / 2365. Взяты данные Excel.'),
]
# если одного человека в разных файлах пишут по-разному: 'как в Excel' -> 'как в рейтинге'
ALIASES = {'Egrinderolls': 'Егор АА 11', 'Влад Сам': 'Coach krotovski',
           'Матвей МС': 'Матвей Пригожий', 'Саша Немощь': 'Немощь'}  # подтверждено пользователем / чатом

FIELDS = {'участие': 'part', 'старт': 'start', 'добор': 'add', 'цена': 'price', 'выход': 'exit', 'итог': 'itog',
          'выигрыш': 'win', 'кредит': 'credit', 'сделка': 'deal', 'фишки для финала': 'finchips',
          'дилерство': 'dealer', 'участие фф': 'ff'}
SKIP_NAMES = {'количество', 'пропустили', 'новые', 'на один раз', 'имя', 'итого', 'всего'}


def num(v):
    return v if isinstance(v, (int, float)) and not isinstance(v, bool) else None


def load_rating():
    ws = openpyxl.load_workbook(RATING, data_only=True)['Рейтинг Игроков']
    events = {}
    e = 1
    while True:
        col = e + 2
        h = ws.cell(1, col).value
        if h is None or not str(h).isdigit(): break
        pts = {ws.cell(r, 2).value: ws.cell(r, col).value for r in range(2, ws.max_row + 1)
               if ws.cell(r, 2).value and isinstance(ws.cell(r, col).value, (int, float))}
        if pts:
            mx = max(pts.values())
            events[e] = {n: dict(pts=p, place=round(mx * mx / (p * p))) for n, p in pts.items()}
        e += 1
    return events


def load_labels():
    ws = openpyxl.load_workbook(HIST, data_only=True)['Итоги']
    res = {}
    for c in range(9, ws.max_column + 1):
        nm = ws.cell(2, c).value
        if not nm: continue
        pod = {}
        for r in range(3, ws.max_row + 1):
            v = ws.cell(r, c).value
            if isinstance(v, str) and v.strip() != '+':
                pod.setdefault(v.strip().upper(), []).append(ws.cell(r, 2).value)
        res[str(nm).strip()] = pod
    return res


def podium_for(labels, ref):
    r = ref.lower()
    for k, v in labels.items():
        if r in k.lower() or k.lower() in r:
            return v
    return {}


def parse_days(ws):
    starts = [(c.column, str(c.value)) for c in ws[4] if c.value and str(c.value).startswith(('День', 'Финал'))]
    days, final_date = [], None
    for i, (col, label) in enumerate(starts):
        if label.startswith('Финал'):
            final_date = label.replace('Финал', '').strip(); continue
        end = starts[i + 1][0] if i + 1 < len(starts) else ws.max_column + 1
        hdr = {}
        for c in range(col, end):
            h = ws.cell(5, c).value
            if h and str(h).strip().lower() in FIELDS:
                hdr.setdefault(FIELDS[str(h).strip().lower()], c)
        m = re.match(r'День\s+(\d+)\s+([\d.]+)', label)
        days.append(dict(n=int(m.group(1)), date=m.group(2).rstrip('.'), cols=hdr))
    return days, final_date


def parse_players(ws, cfg, days, finalists):
    fin_col = ci(cfg['final_col']) if cfg['final_col'] else None
    price_col = ci(cfg['price_col']) if cfg.get('price_col') else None
    out = {}
    for r in range(6, ws.max_row + 1):
        name = ws.cell(r, 2).value
        if not name: continue
        name = str(name).strip()
        if name.lower() in SKIP_NAMES: continue
        name = ALIASES.get(name, name)
        if not (isinstance(ws.cell(r, 1).value, (int, float)) or name in finalists): continue
        if name in out: continue
        rec = dict(days=[], visited=0)
        for d in days:
            g = lambda k: num(ws.cell(r, d['cols'][k]).value) if k in d['cols'] else None
            part, st_, add, ex = g('part'), g('start'), g('add'), g('exit')
            played = bool(part) or bool(st_)
            rec['visited'] += played
            rec['days'].append(dict(n=d['n'], played=played, addon=(bool(add) if add is not None else None) if played else None,
                                    exit=ex if played else None, itog=g('itog') if played else None, price=g('price') if played else None))
        rec['raw'] = num(ws.cell(r, fin_col).value) if fin_col else None
        rec['price'] = num(ws.cell(r, price_col).value) if price_col else None
        out[name] = rec
    return out



def parse_hunt(ws, players, finalists):
    """Таблица нокаутов справа от основной (есть у А-2 и 3-8): имена, суммы выбитых голов, количество."""
    hdr = 6
    cs = next((c for c in range(1, ws.max_column + 1) if ws.cell(hdr, c).value == 'Сумма'), None)
    if not cs: return None
    nc = max((c for c in range(1, cs) if ws.cell(hdr, c).value == 'Имя'), default=None)
    if not nc or nc < 10: return None
    hunters = []
    for r in range(hdr + 1, ws.max_row + 1):
        nm = ws.cell(r, nc).value
        if not nm: break
        nm = ALIASES.get(str(nm).strip(), str(nm).strip())
        kos = [v for v in (num(ws.cell(r, c).value) for c in range(nc + 1, cs)) if v]
        if not kos: continue
        hunters.append(dict(name=nm, kos=kos, count=len(kos), sum=sum(kos), avg=round(sum(kos) / len(kos)), best=max(kos), finalist=nm in finalists))
    hunters.sort(key=lambda h: (-h['count'], -h['sum']))
    busts_col = next((c.column for c in ws[5] if c.value and str(c.value).strip().lower().startswith('выбыван')), None)
    busts = []
    if busts_col:
        for nm, p in players.items():
            pass
    bounty = []
    for nm, p in players.items():
        vals = [d['price'] for d in p['days'] if d.get('price')] + ([p['price']] if p.get('price') else [])
        if vals: bounty.append(dict(name=nm, price=max(vals), finalist=nm in finalists))
    bounty.sort(key=lambda x: -x['price'])
    return dict(hunters=hunters, bounty=bounty[:10], total_ko=sum(h['count'] for h in hunters), total_sum=sum(h['sum'] for h in hunters), busts_col=busts_col)


def rank_desc(values):
    """спортивный ранг: 1 = максимум; одинаковые значения делят место"""
    srt = sorted(values, reverse=True)
    return [srt.index(v) + 1 for v in values]


def spearman(a, b):
    n = len(a)
    if n < 4: return None
    ra, rb = rank_desc([-x for x in a]), rank_desc([-x for x in b])
    ma, mb = sum(ra) / n, sum(rb) / n
    den = (sum((x - ma) ** 2 for x in ra) * sum((y - mb) ** 2 for y in rb)) ** .5
    return None if den == 0 else sum((x - ma) * (y - mb) for x, y in zip(ra, rb)) / den


def main():
    events = load_rating()
    labels = load_labels()
    result = dict(tournaments=[], players={}, quality=[], generated_from=os.path.basename(RATING), events={})
    for cfg in TOURNAMENTS:
        f = glob.glob(os.path.join(ROOT, '2026', cfg['folder'], '*.xlsx'))
        if not f:
            result['quality'].append(f"{cfg['title']}: файл не найден"); continue
        ws = openpyxl.load_workbook(f[0], data_only=True)['Главная']
        days, final_date = parse_days(ws)
        pod = podium_for(labels, cfg['ref'])
        first = (pod.get('ПЕРВОЕ') or pod.get('ПЕРВЫЙ') or [None])[0]
        second = (pod.get('ВТОРОЕ') or pod.get('ВТОРОЙ') or [None])[0]
        fin_event = None
        for e, ev in events.items():
            if first and ev.get(first, {}).get('place') == 1 and max(v['pts'] for v in ev.values()) > 200:
                if second is None or ev.get(second, {}).get('place') == 2:
                    fin_event = e
        if fin_event is None:
            result['quality'].append(f"{cfg['title']}: не найден столбец рейтинга для финала"); continue
        fin = events[fin_event]
        day_events = [fin_event - len(days) + i for i in range(len(days))]
        players = parse_players(ws, cfg, days, set(fin)) if cfg.get('use_table', True) else {}
        trophy = cfg.get('trophy')
        rows = []
        for name, ev in fin.items():
            p = players.get(name)
            if p is None and not cfg.get('use_table', True):
                p = dict(days=[dict(n=d['n'], played=False, addon=None, exit=None, itog=None, price=None) for d in days], visited=None, raw=None, price=None)
            elif p is None:
                result['quality'].append(f"{cfg['title']}: финалист «{name}» не найден в таблице, стек = только базовый бонус")
                p = dict(days=[dict(n=d['n'], played=False, addon=None, exit=None, itog=None, price=None) for d in days], visited=0, raw=None, price=None)
            have_stack = cfg['final_col'] is not None
            stack = None
            if have_stack:
                tro = 0
                if trophy and p['price']:
                    tro = p['price'] * trophy.get(p['visited'], 0)
                stack = cfg['base'] + (p['raw'] or 0) + tro
            dd = []
            for d, e in zip(p['days'], day_events):
                dev = events.get(e, {})
                dp = None
                ref = dev.get(name)
                if ref: dp = ref['place']
                dd.append(dict(n=d['n'], played=d['played'], addon=d['addon'], exit=d['exit'], itog=d['itog'], price=d['price'], place=dp))
            rows.append(dict(name=name, place=ev['place'], points=ev['pts'], visits=p['visited'], stack=stack, days=dd))
        rows.sort(key=lambda r: r['place'])
        if any(r['stack'] is not None for r in rows):
            ranks = rank_desc([r['stack'] for r in rows])
            tot = sum(r['stack'] for r in rows)
            for r, k in zip(rows, ranks):
                r['stack_rank'] = k
                r['delta'] = k - r['place']          # >0: итог лучше, чем по стеку
                r['stack_share'] = r['stack'] / tot
        # данные о днях
        dsum = []
        for d, e in zip(days, day_events):
            played = [(n, pl) for n, pl in players.items() if pl['days'][d['n'] - 1]['played']]
            adds = [pl['days'][d['n'] - 1]['addon'] for n, pl in played if pl['days'][d['n'] - 1]['addon'] is not None]
            dsum.append(dict(n=d['n'], date=d['date'], players=len(played) or len(events.get(e, {})),
                             addon_rate=(sum(adds) / len(adds)) if adds else None, event=e))
        t = dict(nonstandard=bool(cfg.get('nonstandard')), id=cfg['id'], title=cfg['title'], short=cfg['short'], final_date=final_date, flag=cfg.get('flag'),
                 base=cfg['base'], stack_available=any(r['stack'] is not None for r in rows),
                 days=dsum, finalists=len(rows), rating_event=fin_event, rows=rows)
        hunt = parse_hunt(ws, players, set(fin)) if cfg.get('use_table', True) else None
        if hunt:
            if hunt.pop('busts_col', None):
                bc = next(c.column for c in ws[5] if c.value and str(c.value).strip().lower().startswith('выбыван'))
                hunt['busts'] = sorted([dict(name=ALIASES.get(str(ws.cell(r, 2).value).strip(), str(ws.cell(r, 2).value).strip()), n=int(ws.cell(r, bc).value)) for r in range(6, ws.max_row + 1) if ws.cell(r, 2).value and isinstance(ws.cell(r, bc).value, (int, float)) and ws.cell(r, bc).value > 0], key=lambda x: -x['n'])[:10]
            t['hunt'] = hunt
        if cfg.get('note'): t['note'] = cfg['note']
        # метрики турнира
        if t['stack_available']:
            lead = min(rows, key=lambda r: r['stack_rank'])
            t['chip_leader'] = dict(name=lead['name'], place=lead['place'])
            t['winner'] = dict(name=rows[0]['name'], stack_rank=rows[0]['stack_rank'])
            t['corr'] = spearman([-r['stack'] for r in rows], [r['place'] for r in rows])
            t['top3_stack_avg_place'] = st.mean(r['place'] for r in rows if r['stack_rank'] <= 3)
        result['tournaments'].append(t)
        for d in dsum: result['events'][str(d['event'])] = dict(label=f"{cfg['short']} · день {d['n']}", kind='day', tid=cfg['id'])
        result['events'][str(fin_event)] = dict(label=f"{cfg['short']} · финал", kind='final', tid=cfg['id'])
        print(f"{cfg['short']:>4}  финал: событие {fin_event:>2}, финалистов {len(rows):>2}, дней {len(days)}, победитель {rows[0]['name']}", file=sys.stderr)

    # сводка по игрокам
    pl = {}
    for t in result['tournaments']:
        for r in t['rows']:
            p = pl.setdefault(r['name'], dict(name=r['name'], finals=0, wins=0, podiums=0, table=0, places=[], deltas=[], list=[]))
            p['finals'] += 1; p['places'].append(r['place'])
            p['wins'] += r['place'] == 1; p['podiums'] += r['place'] <= 3; p['table'] += r['place'] <= 9
            if r.get('delta') is not None: p['deltas'].append(r['delta'])
            p['list'].append(dict(t=t['id'], short=t['short'], place=r['place'], stack_rank=r.get('stack_rank'), delta=r.get('delta'),
                                  stack=r['stack'], visits=r['visits'], of=t['finalists']))
    for p in pl.values():
        p['avg_place'] = st.mean(p['places'])
        p['sd_place'] = st.pstdev(p['places']) if len(p['places']) > 1 else None
        p['avg_delta'] = st.mean(p['deltas']) if p['deltas'] else None

    # --- рейтинг: история по событиям (лучшие RATING_TOP результатов, как в формуле Excel) ---
    RATING_TOP = 12
    wsr = openpyxl.load_workbook(RATING, data_only=True)['Рейтинг Игроков']
    names_all = [wsr.cell(r, 2).value for r in range(2, wsr.max_row + 1) if wsr.cell(r, 2).value]
    evs = sorted(events)
    running = {n: [] for n in names_all}
    rating_after = {}
    for e in evs:
        for n in names_all:
            if n in events[e]: running[n].append(events[e][n]['pts'])
        rating_after[e] = {n: sum(sorted(running[n], reverse=True)[:RATING_TOP]) for n in names_all}
    rank_after = {}
    for e in evs:
        order = sorted(names_all, key=lambda n: -rating_after[e][n])
        rank_after[e] = {n: i + 1 for i, n in enumerate(order) if rating_after[e][n] > 0}
    for e in evs: result['events'].setdefault(str(e), dict(label=f'Событие {e}', kind='other', tid=None))
    result['rating_top'] = RATING_TOP
    # --- паспорт из листа 'Итоги' (без дней рождения и контактов) ---
    wsi = openpyxl.load_workbook(HIST, data_only=True)['Итоги']
    passport = {}
    for r in range(3, wsi.max_row + 1):
        nm = wsi.cell(r, 2).value
        if not nm: continue
        nm = ALIASES.get(str(nm).strip(), str(nm).strip())
        first = wsi.cell(r, 5).value
        passport[nm] = dict(visits=wsi.cell(r, 4).value if isinstance(wsi.cell(r, 4).value, (int, float)) else None,
                            first=first.strftime('%Y-%m-%d') if hasattr(first, 'strftime') else None,
                            hand=str(wsi.cell(r, 7).value) if wsi.cell(r, 7).value not in (None, '') else None,
                            motto=str(wsi.cell(r, 8).value) if wsi.cell(r, 8).value not in (None, '') else None)
    for n, p in pl.items():
        if n in running:
            p['series'] = [dict(e=e, pts=events[e][n]['pts'], rating=rating_after[e][n], rank=rank_after[e].get(n)) for e in evs if n in events[e]]
            p['rating_now'] = rating_after[evs[-1]][n]; p['rating_rank'] = rank_after[evs[-1]].get(n)
        p['passport'] = passport.get(n)

    # --- проверки качества данных ---
    checks = []
    for q in result['quality']:
        checks.append(dict(level='warn', text=q))
    rated = set(n for e in events.values() for n in e)
    for cfg in TOURNAMENTS:
        t = next((x for x in result['tournaments'] if x['id'] == cfg['id']), None)
        if not t: continue
        if cfg.get('flag'): checks.append(dict(level='info', tid=cfg['id'], text=f"{cfg['title']}: {cfg['flag']}"))
        if cfg.get('note'): checks.append(dict(level='warn', tid=cfg['id'], text=f"{cfg['title']}: {cfg['note']}"))
        f = glob.glob(os.path.join(ROOT, '2026', cfg['folder'], '*.xlsx'))
        if f and cfg.get('use_table', True):
            wsx = openpyxl.load_workbook(f[0], data_only=True)['Главная']
            dd, _ = parse_days(wsx)
            pp = parse_players(wsx, cfg, dd, set())
            fin = {r['name'] for r in t['rows']}
            odd = [n for n, p in pp.items() if p['visited'] and n not in rated and n not in fin]
            if odd:
                import difflib
                parts = []
                for n in odd:
                    cand = difflib.get_close_matches(n, list(rated), 1, 0.6)
                    parts.append(n + (f' (похоже на «{cand[0]}»)' if cand else ''))
                checks.append(dict(level='warn', tid=cfg['id'], text=f"{cfg['title']}: в таблице есть игроки, которых нет в рейтинге: " + ', '.join(parts) + '. Возможно, это другое написание ника — добавьте в ALIASES.'))
    checks.append(dict(level='info', text=f'Рейтинг считается как сумма {RATING_TOP} лучших результатов (как в формуле Excel); места в финалах восстановлены по очкам: место = (макс. очки / очки)².'))
    result['checks'] = checks
    result['players'] = pl
    npath = os.path.join(HERE, 'names.json')
    result['display'] = json.load(open(npath, encoding='utf-8')) if os.path.exists(npath) else {}
    os.makedirs(os.path.join(HERE, 'data'), exist_ok=True)
    with open(os.path.join(HERE, 'data', 'data.json'), 'w', encoding='utf-8') as fh:
        json.dump(result, fh, ensure_ascii=False)
    with open(os.path.join(HERE, 'data', 'data.js'), 'w', encoding='utf-8') as fh:
        fh.write('window.NEPOKER = ' + json.dumps(result, ensure_ascii=False) + ';')
    print('Готово:', len(result['tournaments']), 'турниров,', len(pl), 'игроков; замечаний:', len(result['quality']), file=sys.stderr)
    for q in result['quality']: print(' -', q, file=sys.stderr)


if __name__ == '__main__':
    main()
